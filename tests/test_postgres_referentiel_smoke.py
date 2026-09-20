"""
Smoke test PostgreSQL du référentiel académique.

Ce test ne remplace pas les tests unitaires SQLite : il vérifie le chemin
réel de production (PostgreSQL + Alembic + startup FastAPI + import Excel)
sur une base neuve, puis vérifie que le second passage est idempotent.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, select, func

from app.main import app
from app.models import Domaine, Mention, ProgrammeUniversitaire, Universite, Filiere, Faculte
from scripts.import_academic_data import importer


PERIMETRE_PUBLIC = {
    "universite d'antananarivo",
    "universite d'antsiranana",
    "universite de fianarantsoa",
    "universite de mahajanga",
    "universite de toamasina",
    "universite de toliara",
}


def _normaliser(texte: str | None) -> str:
    import re
    import unicodedata

    if not texte:
        return ""
    texte = texte.strip().replace("\u2019", "'")
    texte = unicodedata.normalize("NFKD", texte)
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texte).lower()


def _source_referentiel() -> tuple[set[str], set[str]]:
    chemin = Path(__file__).resolve().parent.parent / "mahay_universites_mentions_filieres_recensement.xlsx"
    assert chemin.exists(), f"Classeur absent du dépôt : {chemin}"

    classeur = load_workbook(chemin, read_only=True, data_only=True)
    feuille = classeur.active
    lignes = list(feuille.iter_rows(values_only=True))
    entetes = [str(cell).strip() for cell in lignes[0]]
    index = {nom: i for i, nom in enumerate(entetes)}

    domaines: set[str] = set()
    mentions: set[str] = set()
    for ligne in lignes[1:]:
        if not any(ligne):
            continue
        universite = str(ligne[index["Université"]] or "").strip()
        if _normaliser(universite) not in PERIMETRE_PUBLIC:
            continue
        domaine = str(ligne[index["Domaine"]] or "").strip()
        mention = str(ligne[index["Mention"]] or "").strip()
        if domaine:
            domaines.add(domaine)
        if mention:
            mentions.add(mention)
    return domaines, mentions


@pytest.mark.skipif(
    not os.environ.get("DATABASE_URL", "").startswith(("postgresql://", "postgresql+psycopg2://")),
    reason="Smoke test réservé à PostgreSQL",
)
def test_postgres_demarrage_import_et_idempotence():
    # Le TestClient déclenche réellement le startup FastAPI :
    # Alembic, import du référentiel Excel, seed et provisionnement des cercles.
    with TestClient(app) as client:
        response = client.get("/")
        assert response.status_code == 200

    domaines_source, mentions_source = _source_referentiel()

    from app.database import engine

    with Session(engine) as session:
        universites = {u.nom for u in session.exec(select(Universite)).all()}
        domaines = {d.nom for d in session.exec(select(Domaine)).all()}
        mentions = session.exec(select(Mention)).all()

        assert {
            "Universite de Toamasina",
            "Universite d'Antananarivo",
            "Universite de Fianarantsoa",
            "Universite de Mahajanga",
            "Universite de Toliara",
            "Universite d'Antsiranana",
        }.issubset(universites)

        assert domaines_source.issubset(domaines)

        mentions_par_nom = {_normaliser(m.nom): m for m in mentions}
        for nom in mentions_source:
            assert _normaliser(nom) in mentions_par_nom, f"Mention Excel absente de la base : {nom}"
            assert mentions_par_nom[_normaliser(nom)].domaine_id is not None, (
                f"Mention sans Domaine après synchronisation : {nom}"
            )

        uto = next(u for u in session.exec(select(Universite)).all() if _normaliser(u.nom) == "universite de toamasina")
        nb_filieres_uto = session.exec(
            select(func.count())
            .select_from(Filiere)
            .join(Faculte, Filiere.faculte_id == Faculte.id)
            .where(Faculte.universite_id == uto.id)
        ).one()
        nb_programmes_uto = session.exec(
            select(func.count())
            .select_from(ProgrammeUniversitaire)
            .where(ProgrammeUniversitaire.universite_id == uto.id)
        ).one()

        assert nb_filieres_uto == nb_programmes_uto

        rapport_second_passage = importer(
            str(Path(__file__).resolve().parent.parent / "mahay_universites_mentions_filieres_recensement.xlsx")
        )
        assert rapport_second_passage.domaines_crees == []
        assert rapport_second_passage.mentions_creees == []
        assert rapport_second_passage.programmes_crees == 0

        nb_domaines_apres = session.exec(select(func.count()).select_from(Domaine)).one()
        nb_mentions_apres = session.exec(select(func.count()).select_from(Mention)).one()

    assert nb_domaines_apres == len(domaines)
    assert nb_mentions_apres == len(mentions)
